from io import BytesIO
from pptx import Presentation
from pptx.util import Inches

from hackpack.models import Claim, EvidenceType, PitchDeck, SlideContent
from hackpack.generators.pptx_engine import generate_pptx
from hackpack.generators.script_engine import generate_script


def make_deck(with_claims=False):
    claims = [
        Claim(
            statement="The repository contains an app entry point",
            source="app/main.py",
            evidence_type=EvidenceType.REPO_EVIDENCE,
            needs_action="Confirm with a live demo.",
        )
    ] if with_claims else []
    return PitchDeck(
        project_name="Example Project",
        tagline="A simple, honest demo",
        total_duration_seconds=180,
        slides=[
            SlideContent(
                slide_number=i,
                title=title,
                subtitle="A short explanation",
                bullet_points=["First grounded point", "Second grounded point"],
                speaker_notes="Explain the feature and show what works today.",
                speaking_time_seconds=30,
                claims=list(claims),
            )
            for i, title in enumerate(["Problem", "Solution", "How it works", "Demo", "Evidence", "Next steps"], 1)
        ],
    )


def slide_text(slide) -> str:
    return " ".join(
        shape.text for shape in slide.shapes if getattr(shape, "has_text_frame", False)
    )


def test_pptx_is_valid_and_editable():
    payload = generate_pptx(make_deck())
    assert isinstance(payload, bytes)
    prs = Presentation(BytesIO(payload))
    assert len(prs.slides) == 6
    first_slide_text = slide_text(prs.slides[0])
    assert "Problem" in first_slide_text
    assert "First grounded point" in first_slide_text


def test_pptx_dimensions_are_16_by_9():
    prs = Presentation(BytesIO(generate_pptx(make_deck())))
    ratio = prs.slide_width / prs.slide_height
    assert abs(ratio - (16 / 9)) < 0.02
    assert abs(prs.slide_width - Inches(13.333)) < 1000


def test_pptx_background_is_dark():
    prs = Presentation(BytesIO(generate_pptx(make_deck())))
    for slide in prs.slides:
        fill = slide.background.fill
        assert fill.type is not None  # a real solid fill was applied
        rgb = fill.fore_color.rgb
        assert rgb is not None
        # Dark background: each channel below 60.
        assert max(rgb[0], rgb[1], rgb[2]) < 60


def test_pptx_speaker_notes_are_attached():
    prs = Presentation(BytesIO(generate_pptx(make_deck())))
    notes = prs.slides[0].notes_slide.notes_text_frame.text
    assert "Explain the feature" in notes
    assert "HackPack timing" in notes


def test_pptx_notes_include_claim_warnings():
    prs = Presentation(BytesIO(generate_pptx(make_deck(with_claims=True))))
    notes = prs.slides[0].notes_slide.notes_text_frame.text
    assert "Review before presenting" in notes
    assert "Confirm with a live demo." in notes


def test_pptx_evidence_footer_shows_sources():
    prs = Presentation(BytesIO(generate_pptx(make_deck(with_claims=True))))
    text = slide_text(prs.slides[0])
    assert "Evidence" in text
    assert "app/main.py" in text


def test_pptx_saves_to_path(tmp_path):
    target = tmp_path / "deck.pptx"
    result = generate_pptx(make_deck(), target)
    assert result == str(target)
    assert target.exists()
    assert len(Presentation(str(target)).slides) == 6


def test_script_has_timestamps_and_all_slides():
    script = generate_script(make_deck())
    assert "Slide 1: Problem" in script
    assert "Slide 6: Next steps" in script
    assert "Total timed duration" in script


def test_script_lists_evidence_references():
    script = generate_script(make_deck(with_claims=True))
    assert "Evidence references" in script
    assert "`app/main.py`" in script


def test_script_labels_fallback_mode():
    deck = make_deck()
    deck.generation_mode = "fallback"
    deck.generation_warnings = ["Model unreachable; template used."]
    script = generate_script(deck)
    assert "fallback" in script
    assert "Model unreachable" in script
