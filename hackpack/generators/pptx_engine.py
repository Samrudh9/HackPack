"""Render validated pitch content into an editable PowerPoint file."""
from __future__ import annotations
from io import BytesIO
from pathlib import Path
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from hackpack.models import PitchDeck
from hackpack.generators.pptx_template import create_hackpack_template, add_base_slide, add_title, add_bullets, add_evidence_footer, add_accent_panel, CYAN


def _add_speaker_notes(slide, deck: PitchDeck, content) -> None:
    """Attach the slide's speaking notes using PowerPoint's native notes pane."""
    notes_parts = [content.speaker_notes.strip()] if content.speaker_notes.strip() else []
    notes_parts.append(f"[HackPack timing] Slide {content.slide_number} of {len(deck.slides)} — about {content.speaking_time_seconds} seconds.")
    flagged = [claim for claim in content.claims if claim.needs_action]
    if flagged:
        notes_parts.append("[Review before presenting]")
        notes_parts.extend(f"- {claim.statement}: {claim.needs_action}" for claim in flagged[:5])
    try:
        notes_frame = slide.notes_slide.notes_text_frame
        notes_frame.text = "\n".join(notes_parts)
    except Exception:
        # Notes are a convenience; the full timed script is always exported as Markdown.
        pass


def generate_pptx(deck: PitchDeck, output_path: str | Path | None = None) -> bytes | str:
    """Generate a fully editable PPTX. Returns bytes, or saves and returns a path."""
    prs = create_hackpack_template()
    for index, content in enumerate(deck.slides, start=1):
        slide = add_base_slide(prs, index, deck.project_name, content.title)
        add_title(slide, content.title, content.subtitle)
        add_bullets(slide, content.bullet_points)
        if content.claims:
            sources = [
                claim.source for claim in content.claims
                if claim.source and claim.source.lower() not in {"team", "team context", "team interview"}
            ]
        else:
            sources = []
        if sources:
            add_evidence_footer(slide, sources)
        else:
            add_accent_panel(slide)
            # A compact duration cue helps teams rehearse. Full script is exported separately.
            cue = slide.shapes.add_textbox(Inches(0.9), Inches(6.12), Inches(11.2), Inches(0.24))
            cue.text_frame.clear()
            p = cue.text_frame.paragraphs[0]
            p.text = f"Speaking time: about {content.speaking_time_seconds} seconds"
            p.font.name = "Aptos"
            p.font.size = Pt(10)
            p.font.color.rgb = RGBColor(0, 217, 255)
        _add_speaker_notes(slide, deck, content)
    if output_path is not None:
        target = Path(output_path)
        target.parent.mkdir(parents=True, exist_ok=True)
        prs.save(str(target))
        return str(target)
    buffer = BytesIO()
    prs.save(buffer)
    return buffer.getvalue()
